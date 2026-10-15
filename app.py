# importações de dados.

import os
import json
import sqlite3
import re
import tempfile
from datetime import datetime
from flask import Flask, render_template_string, request, redirect, url_for, flash, send_file, session
from werkzeug.security import generate_password_hash, check_password_hash
import openpyxl
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.chart import BarChart, Reference
import io
from pypdf import PdfReader

app = Flask(__name__)
app.secret_key = "segredo_residuos_2026_sessao"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

CAMINHO_USUARIOS = os.path.join(BASE_DIR, "usuarios.json")
CAMINHO_DB = os.path.join(BASE_DIR, "dados_residuos.db")

usuarios_padrao = {
    "admin": {
        "hash": generate_password_hash("ebserh2026"),
        "role": "admin"
    },
    "suporte": {
        "hash": generate_password_hash("hubrasil2026"),
        "role": "admin"
    }
}

def carregar_usuarios():
    if not os.path.exists(CAMINHO_USUARIOS):
        salvar_usuarios(usuarios_padrao)
        return usuarios_padrao
    with open(CAMINHO_USUARIOS, "r", encoding="utf-8") as f:
        return json.load(f)

def salvar_usuarios(usuarios):
    with open(CAMINHO_USUARIOS, "w", encoding="utf-8") as f:
        json.dump(usuarios, f, indent=4)

def usuario_admin():
    user = session.get("usuario")
    if not user:
        return False
    if user in ["admin", "suporte"]:
        return True
    usuarios = carregar_usuarios()
    return usuarios.get(user, {}).get("role", "").lower() == "admin"

def garantir_banco():
    conn = sqlite3.connect(CAMINHO_DB)
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS residuos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ano TEXT NOT NULL,
            mes TEXT NOT NULL,
            data TEXT NOT NULL,
            localidade TEXT NOT NULL,
            qtd_infectante INTEGER DEFAULT 0,
            peso_infectante REAL DEFAULT 0.0,
            qtd_organico INTEGER DEFAULT 0,
            peso_organico REAL DEFAULT 0.0,
            qtd_comum INTEGER DEFAULT 0,
            peso_comum REAL DEFAULT 0.0,
            qtd_reciclavel INTEGER DEFAULT 0,
            peso_reciclavel REAL DEFAULT 0.0,
            qtd_quimico INTEGER DEFAULT 0,
            peso_quimico REAL DEFAULT 0.0
        )
    """)
    conn.commit()
    conn.close()

# -------------------------------------------------------------------
# Interface do sistema em HTML.
# -------------------------------------------------------------------
HTML_LOGIN = """
<!DOCTYPE html>
<html lang="pt-br">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Login | Gestão de Resíduos</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
    <style> body { font-family: 'Inter', sans-serif; } </style>
</head>
<body class="bg-slate-900 min-h-screen w-screen flex items-center justify-center p-4">

    <div class="w-full max-w-md bg-white rounded-2xl shadow-2xl overflow-hidden border border-slate-700 p-8 space-y-6">
        <div class="text-center space-y-2">
            <div class="inline-flex p-3.5 bg-indigo-50 rounded-2xl text-indigo-600 mb-1">
                <i data-lucide="shield-check" class="w-8 h-8"></i>
            </div>
            <h1 class="text-2xl font-bold text-slate-800">Gestão de Resíduos</h1>
            <p class="text-xs text-slate-500">Introduza as suas credenciais para aceder ao sistema</p>
        </div>

        {% with messages = get_flashed_messages(with_categories=true) %}
          {% if messages %}
            {% for category, message in messages %}
            <div class="p-3.5 {% if category == 'error' %}bg-rose-50 border-rose-200 text-rose-800{% else %}bg-emerald-50 border-emerald-200 text-emerald-800{% endif %} border rounded-xl flex items-center gap-3 text-xs font-medium">
                <i data-lucide="{% if category == 'error' %}alert-circle{% else %}check-circle-2{% endif %}" class="w-4 h-4 shrink-0"></i>
                <span>{{ message }}</span>
            </div>
            {% endfor %}
          {% endif %}
        {% endwith %}

        <form action="/login" method="POST" class="space-y-4">
            <div>
                <label class="block text-xs font-bold text-slate-600 uppercase mb-1">Utilizador</label>
                <input type="text" name="username" required class="w-full px-3.5 py-3 bg-slate-50 border border-slate-300 rounded-xl text-sm focus:ring-2 focus:ring-indigo-600 focus:outline-none">
            </div>
            <div>
                <label class="block text-xs font-bold text-slate-600 uppercase mb-1">Palavra-passe</label>
                <input type="password" name="password" required class="w-full px-3.5 py-3 bg-slate-50 border border-slate-300 rounded-xl text-sm focus:ring-2 focus:ring-indigo-600 focus:outline-none">
            </div>
            <button type="submit" class="w-full py-3.5 px-4 bg-indigo-600 hover:bg-indigo-700 text-white font-semibold rounded-xl shadow-lg transition flex items-center justify-center gap-2 text-sm cursor-pointer">
                <i data-lucide="log-in" class="w-4 h-4"></i> Entrar no Sistema
            </button>
        </form>
    </div>

    <script> lucide.createIcons(); </script>
</body>
</html>
"""

HTML_FORMULARIO = """
<!DOCTYPE html>
<html lang="pt-br">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Lançamento de Resíduos | 2026</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
    <style> body { font-family: 'Inter', sans-serif; } </style>
</head>
<body class="bg-slate-100 min-h-screen w-full m-0 p-4 xl:p-6 flex flex-col justify-between">

    <header class="w-full bg-gradient-to-r from-blue-900 via-indigo-900 to-violet-900 rounded-2xl shadow-xl p-6 text-white mb-6 flex flex-col md:flex-row items-center justify-between gap-4">
        <div class="flex items-center gap-4">
            <img src="{{ url_for('static', filename='logo_direita.jpg') }}" alt="Logo Direita" class="h-14 w-auto object-contain bg-white/10 p-1.5 rounded-xl">
            <div>
                <h1 class="text-2xl font-bold tracking-tight">Gestão Inteligente de Resíduos Hospitalares</h1>
                <p class="text-blue-200 text-xs mt-0.5">Sistema de Lançamento, Classificação e Monitoramento de Pesagem (2026)</p>
            </div>
        </div>

        <div class="flex items-center gap-3">
            {% if e_admin %}
            <a href="/usuarios" title="Gerenciar Usuários" class="flex items-center gap-1.5 bg-white/20 hover:bg-white/30 px-3.5 py-2 rounded-xl text-xs font-semibold backdrop-blur-sm transition">
                <i data-lucide="users" class="w-4 h-4 text-yellow-300"></i>
                <span>Usuários</span>
            </a>
            <a href="/registros" title="Gerenciar Registros" class="flex items-center gap-1.5 bg-white/20 hover:bg-white/30 px-3.5 py-2 rounded-xl text-xs font-semibold backdrop-blur-sm transition">
                <i data-lucide="list-checks" class="w-4 h-4 text-emerald-300"></i>
                <span>Registros</span>
            </a>
            {% endif %}
            
            <div class="flex items-center gap-2 bg-white/10 px-4 py-2 rounded-xl text-xs font-semibold border border-white/10">
                <i data-lucide="user" class="w-4 h-4 text-emerald-400"></i>
                <span>{{ usuario_atual }}</span>
            </div>

            <a href="/logout" title="Sair" class="flex items-center justify-center bg-rose-500 hover:bg-rose-600 p-2 rounded-xl text-xs font-medium transition shadow">
                <i data-lucide="log-out" class="w-4 h-4 text-white"></i>
            </a>
        </div>
    </header>

    {% with messages = get_flashed_messages(with_categories=true) %}
      {% if messages %}
        {% for category, message in messages %}
        <div class="mb-6 p-4 {% if category == 'error' %}bg-rose-50 border-rose-200 text-rose-800{% else %}bg-emerald-50 border-emerald-200 text-emerald-800{% endif %} border rounded-2xl flex items-center gap-3 text-sm font-medium shadow-sm">
            <i data-lucide="{% if category == 'error' %}alert-circle{% else %}check-circle-2{% endif %}" class="w-5 h-5 shrink-0"></i>
            <span>{{ message }}</span>
        </div>
        {% endfor %}
      {% endif %}
    {% endwith %}

    <main class="w-full grid grid-cols-1 xl:grid-cols-12 gap-6 items-start flex-1">
        
        <div class="w-full xl:col-span-5 bg-white rounded-2xl shadow-xl overflow-hidden border border-slate-200/80 p-6 space-y-5">
            <h2 class="text-lg font-bold text-slate-800 flex items-center gap-2 border-b pb-3">
                <i data-lucide="plus-circle" class="w-5 h-5 text-emerald-600"></i> Novo Lançamento de Resíduos
            </h2>

            <form action="/salvar" method="POST" class="space-y-4">
                <div class="grid grid-cols-3 gap-3">
                    <div>
                        <label class="block text-xs font-bold text-slate-600 uppercase mb-1">Data</label>
                        <input type="date" name="data" id="input_data" required class="w-full px-3 py-2.5 bg-slate-50 border border-slate-300 rounded-xl text-xs font-medium">
                    </div>
                    <div>
                        <label class="block text-xs font-bold text-slate-600 uppercase mb-1">Mês Ref.</label>
                        <select name="mes" id="select_mes" required class="w-full px-3 py-2.5 bg-slate-50 border border-slate-300 rounded-xl text-xs font-medium">
                            <option value="JANEIRO">JANEIRO</option>
                            <option value="FEVEREIRO">FEVEREIRO</option>
                            <option value="MARÇO">MARÇO</option>
                            <option value="ABRIL">ABRIL</option>
                            <option value="MAIO">MAIO</option>
                            <option value="JUNHO">JUNHO</option>
                            <option value="JULHO">JULHO</option>
                            <option value="AGOSTO">AGOSTO</option>
                            <option value="SETEMBRO">SETEMBRO</option>
                            <option value="OUTUBRO">OUTUBRO</option>
                            <option value="NOVEMBRO">NOVEMBRO</option>
                            <option value="DEZEMBRO">DEZEMBRO</option>
                        </select>
                    </div>
                    <div>
                        <label class="block text-xs font-bold text-slate-600 uppercase mb-1">Ano Ref.</label>
                        <input type="number" name="ano" id="input_ano" min="2000" max="2100" required class="w-full px-3 py-2.5 bg-slate-50 border border-slate-300 rounded-xl text-xs font-medium">
                    </div>
                </div>

                <div>
                    <label class="block text-xs font-bold text-slate-600 uppercase mb-1">Localidade / Setor</label>
                    <select name="localidade" required class="w-full px-3.5 py-2.5 bg-slate-50 border border-slate-300 rounded-xl text-xs font-medium">
                        <option value="" disabled selected>Selecione o setor...</option>
                        <option value="ANEXO - B">ANEXO - B</option>
                        <option value="ANEXO - H">ANEXO - H</option>
                        <option value="CASA DA ALERGIA">CASA DA ALERGIA</option>
                        <option value="CASA NEFRO">CASA NEFRO</option>
                        <option value="CENTRAL">CENTRAL</option>
                        <option value="LABORATÓRIO">LABORATÓRIO</option>
                        <option value="MAMOGRAFIA">MAMOGRAFIA</option>
                        <option value="MATERNIDADE">MATERNIDADE</option>
                        <option value="NUTRIÇÃO">NUTRIÇÃO</option>
                    </select>
                </div>

                <div class="grid grid-cols-1 md:grid-cols-2 gap-3 max-h-[360px] overflow-y-auto pr-1">
                    <div class="bg-rose-50/50 p-3 rounded-xl border border-rose-100 space-y-2">
                        <span class="text-xs font-bold text-rose-700 uppercase tracking-wider block">Infectante</span>
                        <div class="grid grid-cols-2 gap-2">
                            <div><label class="block text-[10px] font-semibold text-slate-500 mb-0.5">Qtd. Carrinhos</label><input type="number" name="qtd_infectante" value="0" min="0" required class="w-full px-2 py-1.5 bg-white border rounded-lg text-xs"></div>
                            <div><label class="block text-[10px] font-semibold text-slate-500 mb-0.5">Peso (Kg)</label><input type="number" name="peso_infectante" value="0.0" min="0" step="0.01" required class="w-full px-2 py-1.5 bg-white border rounded-lg text-xs"></div>
                        </div>
                    </div>

                    <div class="bg-amber-50/50 p-3 rounded-xl border border-amber-100 space-y-2">
                        <span class="text-xs font-bold text-amber-700 uppercase tracking-wider block">Compostagem Orgânico</span>
                        <div class="grid grid-cols-2 gap-2">
                            <div><label class="block text-[10px] font-semibold text-slate-500 mb-0.5">Qtd. Carrinhos</label><input type="number" name="qtd_organico" value="0" min="0" required class="w-full px-2 py-1.5 bg-white border rounded-lg text-xs"></div>
                            <div><label class="block text-[10px] font-semibold text-slate-500 mb-0.5">Peso (Kg)</label><input type="number" name="peso_organico" value="0.0" min="0" step="0.01" required class="w-full px-2 py-1.5 bg-white border rounded-lg text-xs"></div>
                        </div>
                    </div>

                    <div class="bg-slate-50 p-3 rounded-xl border border-slate-200/60 space-y-2">
                        <span class="text-xs font-bold text-slate-600 uppercase tracking-wider block">Comum Geral</span>
                        <div class="grid grid-cols-2 gap-2">
                            <div><label class="block text-[10px] font-semibold text-slate-500 mb-0.5">Qtd. Carrinhos</label><input type="number" name="qtd_comum" value="0" min="0" required class="w-full px-2 py-1.5 bg-white border rounded-lg text-xs"></div>
                            <div><label class="block text-[10px] font-semibold text-slate-500 mb-0.5">Peso (Kg)</label><input type="number" name="peso_comum" value="0.0" min="0" step="0.01" required class="w-full px-2 py-1.5 bg-white border rounded-lg text-xs"></div>
                        </div>
                    </div>

                    <div class="bg-emerald-50/50 p-3 rounded-xl border border-emerald-100 space-y-2">
                        <span class="text-xs font-bold text-emerald-700 uppercase tracking-wider block">Reciclagem Cedo</span>
                        <div class="grid grid-cols-2 gap-2">
                            <div><label class="block text-[10px] font-semibold text-slate-500 mb-0.5">Qtd. Carrinhos</label><input type="number" name="qtd_reciclavel" value="0" min="0" required class="w-full px-2 py-1.5 bg-white border rounded-lg text-xs"></div>
                            <div><label class="block text-[10px] font-semibold text-slate-500 mb-0.5">Peso (Kg)</label><input type="number" name="peso_reciclavel" value="0.0" min="0" step="0.01" required class="w-full px-2 py-1.5 bg-white border rounded-lg text-xs"></div>
                        </div>
                    </div>

                    <div class="bg-purple-50/50 p-3 rounded-xl border border-purple-100 space-y-2 md:col-span-2">
                        <span class="text-xs font-bold text-purple-700 uppercase tracking-wider block">Químico</span>
                        <div class="grid grid-cols-2 gap-2">
                            <div><label class="block text-[10px] font-semibold text-slate-500 mb-0.5">Qtd. Carrinhos</label><input type="number" name="qtd_quimico" value="0" min="0" required class="w-full px-2 py-1.5 bg-white border rounded-lg text-xs"></div>
                            <div><label class="block text-[10px] font-semibold text-slate-500 mb-0.5">Peso (Kg)</label><input type="number" name="peso_quimico" value="0.0" min="0" step="0.01" required class="w-full px-2 py-1.5 bg-white border rounded-lg text-xs"></div>
                        </div>
                    </div>
                </div>

                <button type="submit" class="w-full py-3 px-4 bg-emerald-600 hover:bg-emerald-700 text-white font-semibold rounded-xl shadow-lg transition flex items-center justify-center gap-2 cursor-pointer text-sm">
                    <i data-lucide="save" class="w-4 h-4"></i> Salvar Lançamento Manual
                </button>
            </form>

            <hr class="border-slate-100">

            <form action="/importar/pdf" method="POST" enctype="multipart/form-data" class="space-y-2">
                <label class="block text-xs font-bold text-slate-600 uppercase tracking-wider">Importar Relatório em PDF</label>
                <div class="flex items-center gap-2">
                    <input type="file" name="arquivo_pdf" accept=".pdf" required class="w-full text-xs text-slate-500 file:mr-3 file:py-2 file:px-3 file:rounded-xl file:border-0 file:text-xs file:font-semibold file:bg-violet-50 file:text-violet-700 hover:file:bg-violet-100 cursor-pointer">
                    <button type="submit" class="py-2.5 px-4 bg-violet-600 hover:bg-violet-700 text-white font-semibold rounded-xl shadow text-xs transition shrink-0 cursor-pointer flex items-center gap-1">
                        <i data-lucide="upload" class="w-3.5 h-3.5"></i> Importar
                    </button>
                </div>
            </form>

            <a href="/relatorio/filtros" class="w-full py-3 px-4 bg-indigo-600 hover:bg-indigo-700 text-white font-semibold rounded-xl shadow transition flex items-center justify-center gap-2 text-sm">
                <i data-lucide="file-spreadsheet" class="w-4 h-4 text-yellow-300"></i> Gerar Relatório em Excel (Com Filtros)
            </a>
        </div>

        <div class="w-full xl:col-span-7 bg-white rounded-2xl shadow-xl p-6 border border-slate-200/80 flex flex-col justify-between min-h-[580px]">
            <div>
                <div class="flex items-center justify-between mb-4 border-b pb-3">
                    <div>
                        <h3 class="text-lg font-bold text-slate-800 flex items-center gap-2">
                            <i data-lucide="bar-chart-3" class="w-5 h-5 text-indigo-600"></i> Comparativo Mensal Consolidado (Kg)
                        </h3>
                        <p class="text-xs text-slate-500">Infectante, Compostagem, Comum, Reciclagem e Químico</p>
                    </div>
                    
                    <button type="button" onclick="abrirModalGrafico()" class="py-2 px-3.5 bg-indigo-50 hover:bg-indigo-100 text-indigo-700 text-xs font-bold rounded-xl border border-indigo-200 flex items-center gap-1.5 transition cursor-pointer">
                        <i data-lucide="maximize-2" class="w-4 h-4"></i> Expandir / Baixar Gráfico
                    </button>
                </div>

                <div class="relative w-full h-[460px]">
                    <canvas id="graficoResiduos"></canvas>
                </div>
            </div>
        </div>

    </main>

    <div id="modalGrafico" class="fixed inset-0 bg-slate-900/75 backdrop-blur-sm z-50 hidden flex items-center justify-center p-4">
        <div class="bg-white rounded-2xl shadow-2xl w-full max-w-6xl p-6 relative border border-slate-200 flex flex-col">
            <div class="flex items-center justify-between pb-4 mb-4 border-b">
                <div>
                    <h3 class="text-xl font-bold text-slate-800 flex items-center gap-2">
                        <i data-lucide="bar-chart-2" class="w-6 h-6 text-indigo-600"></i> Visão Detalhada do Gráfico Comparativo
                    </h3>
                    <p class="text-xs text-slate-500 mt-0.5">Indicadores mensais consolidados por categoria de resíduos</p>
                </div>
                
                <div class="flex items-center gap-2">
                    <button onclick="baixarGraficoPNG()" class="py-2 px-4 bg-emerald-600 hover:bg-emerald-700 text-white font-semibold rounded-xl text-xs flex items-center gap-1.5 transition shadow cursor-pointer">
                        <i data-lucide="download" class="w-4 h-4"></i> Baixar Gráfico (.png)
                    </button>
                    <button onclick="fecharModalGrafico()" class="p-2 text-slate-400 hover:text-slate-600 hover:bg-slate-100 rounded-full transition cursor-pointer">
                        <i data-lucide="x" class="w-6 h-6"></i>
                    </button>
                </div>
            </div>
            
            <div class="relative w-full h-[540px]">
                <canvas id="graficoResiduosModal"></canvas>
            </div>
        </div>
    </div>

    <script>
        lucide.createIcons();
        let chartModalInstance = null;

        document.addEventListener("DOMContentLoaded", function() {
            const dateInput = document.querySelector('input[type="date"]');
            const mesSelect = document.getElementById('select_mes');
            const anoInput = document.getElementById('input_ano');

            const hoje = new Date();
            const ano = hoje.getFullYear();
            const mesIndex = hoje.getMonth();
            const dia = String(hoje.getDate()).padStart(2, '0');
            const mesStr = String(mesIndex + 1).padStart(2, '0');
            
            if (dateInput) dateInput.value = `${ano}-${mesStr}-${dia}`;
            if (anoInput) anoInput.value = ano;

            const mesesNomes = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"];
            if (mesSelect) mesSelect.value = mesesNomes[mesIndex];

            if (dateInput) {
                dateInput.addEventListener('change', function() {
                    if (dateInput.value) {
                        const novaData = new Date(dateInput.value + 'T00:00:00');
                        if (!isNaN(novaData.getTime())) {
                            if (anoInput) anoInput.value = novaData.getFullYear();
                            if (mesSelect) mesSelect.value = mesesNomes[novaData.getMonth()];
                        }
                    }
                });
            }

            const ctx = document.getElementById('graficoResiduos').getContext('2d');
            const dg = {{ dados_grafico | tojson }};
            const labels = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"];

            new Chart(ctx, {
                type: 'bar',
                data: {
                    labels: labels,
                    datasets: [
                        { label: 'Infectante', data: labels.map(m => dg.infectante[m] || 0), backgroundColor: 'rgba(225, 29, 72, 0.85)', borderRadius: 4 },
                        { label: 'Compostagem', data: labels.map(m => dg.organico[m] || 0), backgroundColor: 'rgba(217, 119, 6, 0.85)', borderRadius: 4 },
                        { label: 'Comum Geral', data: labels.map(m => dg.comum[m] || 0), backgroundColor: 'rgba(30, 58, 138, 0.85)', borderRadius: 4 },
                        { label: 'Reciclagem', data: labels.map(m => dg.reciclavel[m] || 0), backgroundColor: 'rgba(16, 185, 129, 0.85)', borderRadius: 4 },
                        { label: 'Químico', data: labels.map(m => dg.quimico[m] || 0), backgroundColor: 'rgba(147, 51, 234, 0.85)', borderRadius: 4 }
                    ]
                },
                options: {
                    responsive: true,
                    maintainAspectRatio: false,
                    plugins: { legend: { position: 'top', labels: { font: { family: 'Inter', size: 11 } } } },
                    scales: { y: { beginAtZero: true, grid: { color: 'rgba(226, 232, 240, 0.6)' } }, x: { grid: { display: false } } }
                }
            });
        });

        function abrirModalGrafico() {
            document.getElementById('modalGrafico').classList.remove('hidden');
            const ctxModal = document.getElementById('graficoResiduosModal').getContext('2d');
            const dg = {{ dados_grafico | tojson }};
            const labels = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"];

            if (chartModalInstance) chartModalInstance.destroy();

            chartModalInstance = new Chart(ctxModal, {
                type: 'bar',
                data: {
                    labels: labels,
                    datasets: [
                        { label: 'Infectante', data: labels.map(m => dg.infectante[m] || 0), backgroundColor: 'rgba(225, 29, 72, 0.85)', borderRadius: 6 },
                        { label: 'Compostagem', data: labels.map(m => dg.organico[m] || 0), backgroundColor: 'rgba(217, 119, 6, 0.85)', borderRadius: 6 },
                        { label: 'Comum Geral', data: labels.map(m => dg.comum[m] || 0), backgroundColor: 'rgba(30, 58, 138, 0.85)', borderRadius: 6 },
                        { label: 'Reciclagem', data: labels.map(m => dg.reciclavel[m] || 0), backgroundColor: 'rgba(16, 185, 129, 0.85)', borderRadius: 6 },
                        { label: 'Químico', data: labels.map(m => dg.quimico[m] || 0), backgroundColor: 'rgba(147, 51, 234, 0.85)', borderRadius: 6 }
                    ]
                },
                options: {
                    responsive: true,
                    maintainAspectRatio: false,
                    plugins: { legend: { position: 'top', labels: { font: { family: 'Inter', size: 13 } } } },
                    scales: { y: { beginAtZero: true, grid: { color: 'rgba(226, 232, 240, 0.6)' } }, x: { grid: { display: false } } }
                }
            });
        }

        function fecharModalGrafico() {
            document.getElementById('modalGrafico').classList.add('hidden');
        }

        function baixarGraficoPNG() {
            const canvas = document.getElementById('graficoResiduosModal');
            const imageURL = canvas.toDataURL('image/png');
            const link = document.createElement('a');
            link.href = imageURL;
            link.download = 'grafico_comparativo_residuos.png';
            document.body.appendChild(link);
            link.click();
            document.body.removeChild(link);
        }
    </script>
</body>
</html>
"""

HTML_FILTROS_RELATORIO = """
<!DOCTYPE html>
<html lang="pt-br">
<head>
    <meta charset="UTF-8">
    <title>Filtrar Relatório Excel</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
    <style> body { font-family: 'Inter', sans-serif; } </style>
</head>
<body class="bg-slate-100 min-h-screen flex items-center justify-center p-4">
    <div class="w-full max-w-md bg-white rounded-2xl shadow-xl overflow-hidden border border-slate-200">
        <div class="bg-slate-800 p-6 text-white flex justify-between items-center">
            <h2 class="text-lg font-bold flex items-center gap-2">
                <i data-lucide="filter" class="w-5 h-5 text-yellow-400"></i> Filtros de Relatório Excel
            </h2>
            <a href="/" class="px-3 py-1.5 bg-slate-700 hover:bg-slate-600 text-xs font-semibold rounded-lg flex items-center gap-1">
                <i data-lucide="arrow-left" class="w-4 h-4"></i> Voltar
            </a>
        </div>
        <form action="/relatorio/excel/gerar" method="POST" class="p-6 space-y-4">
            <div>
                <label class="block text-xs font-semibold text-slate-600 uppercase mb-1">Filtrar por Setor</label>
                <select name="filtro_localidade" class="w-full px-3.5 py-2.5 bg-slate-50 border border-slate-300 rounded-lg text-sm">
                    <option value="TODOS">-- Todos os Setores --</option>
                    <option value="ANEXO - B">ANEXO - B</option>
                    <option value="ANEXO - H">ANEXO - H</option>
                    <option value="CASA DA ALERGIA">CASA DA ALERGIA</option>
                    <option value="CASA NEFRO">CASA NEFRO</option>
                    <option value="CENTRAL">CENTRAL</option>
                    <option value="LABORATÓRIO">LABORATÓRIO</option>
                    <option value="MAMOGRAFIA">MAMOGRAFIA</option>
                    <option value="MATERNIDADE">MATERNIDADE</option>
                    <option value="NUTRIÇÃO">NUTRIÇÃO</option>
                </select>
            </div>
            <div>
                <label class="block text-xs font-semibold text-slate-600 uppercase mb-1">Filtrar por Ano</label>
                <input type="text" name="filtro_ano" value="TODOS" class="w-full px-3.5 py-2.5 bg-slate-50 border border-slate-300 rounded-lg text-sm">
            </div>
            <button type="submit" class="w-full py-3 bg-indigo-600 hover:bg-indigo-700 text-white font-semibold rounded-xl shadow text-sm cursor-pointer">
                Baixar Planilha Completa (.xlsx)
            </button>
        </form>
    </div>
    <script> lucide.createIcons(); </script>
</body>
</html>
"""
# usuarios 
HTML_USUARIOS = """
<!DOCTYPE html>
<html lang="pt-br">
<head>
    <meta charset="UTF-8"><title>Usuários</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
    <style> body { font-family: 'Inter', sans-serif; } </style>
</head>
<body class="bg-slate-100 min-h-screen flex items-center justify-center p-4">
    <div class="w-full max-w-xl bg-white rounded-2xl shadow-xl overflow-hidden border">
        <div class="bg-slate-800 p-6 text-white flex justify-between items-center">
            <h2 class="text-xl font-bold flex items-center gap-2"><i data-lucide="users" class="w-5 h-5 text-yellow-400"></i> Usuários</h2>
            <a href="/" class="px-3 py-1.5 bg-slate-700 text-xs rounded-lg flex items-center gap-1"><i data-lucide="arrow-left" class="w-4 h-4"></i> Voltar</a>
        </div>
        
        {% with messages = get_flashed_messages(with_categories=true) %}
          {% if messages %}
            {% for category, message in messages %}
            <div class="m-6 mb-0 p-3 {% if category == 'error' %}bg-rose-50 text-rose-800 border-rose-200{% else %}bg-emerald-50 text-emerald-800 border-emerald-200{% endif %} border rounded-lg text-xs font-medium">
                {{ message }}
            </div>
            {% endfor %}
          {% endif %}
        {% endwith %}

        <form action="/usuarios/criar" method="POST" class="p-6 bg-slate-50 border-b grid grid-cols-1 md:grid-cols-3 gap-3">
            <input type="text" name="novo_usuario" placeholder="Nome de Usuário" required class="px-3 py-2 bg-white border rounded-lg text-sm">
            <input type="password" name="nova_senha" placeholder="Senha" required class="px-3 py-2 bg-white border rounded-lg text-sm">
            <button type="submit" class="py-2 bg-indigo-600 hover:bg-indigo-700 text-white text-xs rounded-lg font-semibold transition">Criar Usuário</button>
        </form>
        <div class="p-6 divide-y max-h-80 overflow-y-auto">
            {% for user, info in lista_usuarios.items() %}
            <div class="py-3 flex items-center justify-between">
                <div><span class="text-sm font-semibold">{{ user }}</span> <span class="text-xs text-slate-400">({{ info.role }})</span></div>
                {% if user not in ['admin', 'suporte'] and user != session.get('usuario') %}
                <form action="/usuarios/deletar" method="POST" onsubmit="return confirm('Tem certeza que deseja remover o usuário {{ user }}?');">
                    <input type="hidden" name="usuario_deletar" value="{{ user }}">
                    <button type="submit" class="text-rose-500 hover:text-rose-700 text-xs font-medium transition">Excluir</button>
                </form>
                {% else %}
                <span class="text-xs text-slate-400 italic">Protegido</span>
                {% endif %}
            </div>
            {% endfor %}
        </div>
    </div>
    <script> lucide.createIcons(); </script>
</body>
</html>
"""

HTML_REGISTROS = """
<!DOCTYPE html>
<html lang="pt-br">
<head>
    <meta charset="UTF-8">
    <title>Gerenciar Registros</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
    <style> body { font-family: 'Inter', sans-serif; } </style>
</head>
<body class="bg-slate-100 min-h-screen p-6">

<div class="max-w-7xl mx-auto bg-white rounded-xl shadow-lg overflow-hidden border">
    <div class="bg-slate-800 text-white p-5 flex justify-between items-center">
        <h1 class="font-bold text-xl flex items-center gap-2">
            <i data-lucide="list-checks" class="w-6 h-6 text-emerald-400"></i> Gerenciar Registros
        </h1>
        <a href="/" class="bg-slate-700 hover:bg-slate-600 px-4 py-2 text-xs font-semibold rounded-lg flex items-center gap-1.5 transition">
            <i data-lucide="arrow-left" class="w-4 h-4"></i> Voltar
        </a>
    </div>

    <div class="p-6 overflow-x-auto">
        <table class="w-full text-left border-collapse text-sm">
            <thead class="bg-slate-100 text-slate-600 uppercase text-xs border-b">
                <tr>
                    <th class="p-3 text-center">ID</th>
                    <th class="p-3 text-center">Ano</th>
                    <th class="p-3 text-center">Mês</th>
                    <th class="p-3 text-center">Data</th>
                    <th class="p-3">Setor</th>
                    <th class="p-3 text-center">Ação</th>
                </tr>
            </thead>
            <tbody class="divide-y divide-slate-200">
                {% for item in registros %}
                <tr class="hover:bg-slate-50 transition">
                    <td class="p-3 text-center font-medium text-slate-500">{{ item.id }}</td>
                    <td class="p-3 text-center font-semibold">{{ item.ano }}</td>
                    <td class="p-3 text-center">{{ item.mes }}</td>
                    <td class="p-3 text-center">{{ item.data }}</td>
                    <td class="p-3 font-medium text-slate-700">{{ item.localidade }}</td>
                    <td class="p-3 text-center">
                        <form action="/registro/excluir" method="POST" onsubmit="return confirm('Tem certeza que deseja excluir o registo ID {{ item.id }}?');">
                            <input type="hidden" name="registro_id" value="{{ item.id }}">
                            <button type="submit" class="bg-rose-50 hover:bg-rose-100 text-rose-600 px-3 py-1.5 rounded-lg text-xs font-semibold inline-flex items-center gap-1 transition">
                                <i data-lucide="trash-2" class="w-3.5 h-3.5"></i> Excluir
                            </button>
                        </form>
                    </td>
                </tr>
                {% endfor %}
            </tbody>
        </table>
    </div>
</div>
<script> lucide.createIcons(); </script>
</body>
</html>
""" 

# -------------------------------------------------------------------
# ROTAS FLASK (Com Autenticação por Sessão)
# -------------------------------------------------------------------
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username").strip().lower()
        password = request.form.get("password").strip()
        
        usuarios = carregar_usuarios()
        if username in usuarios and check_password_hash(usuarios[username]["hash"], password):
            session["usuario"] = username
            return redirect(url_for("index"))
        else:
            flash("Utilizador ou palavra-passe incorretos.", "error")
            
    return render_template_string(HTML_LOGIN)

@app.route("/logout")
def logout():
    session.pop("usuario", None)
    flash("Sessão encerrada com sucesso.", "success")
    return redirect(url_for("login"))

@app.route("/")
def index():
    if "usuario" not in session:
        return redirect(url_for("login"))
        
    usuario_logado = session.get("usuario")
    e_admin = usuario_admin()

    garantir_banco()
    conn = sqlite3.connect(CAMINHO_DB)
    cursor = conn.cursor()
    cursor.execute("""
        SELECT mes, 
               SUM(peso_infectante), 
               SUM(peso_organico), 
               SUM(peso_comum), 
               SUM(peso_reciclavel), 
               SUM(peso_quimico) 
        FROM residuos GROUP BY mes
    """)
    rows = cursor.fetchall()
    conn.close()

    dados_grafico = {
        "infectante": {},
        "organico": {},
        "comum": {},
        "reciclavel": {},
        "quimico": {}
    }
    for mes, p_inf, p_org, p_com, p_rec, p_qui in rows:
        dados_grafico["infectante"][mes] = p_inf or 0.0
        dados_grafico["organico"][mes] = p_org or 0.0
        dados_grafico["comum"][mes] = p_com or 0.0
        dados_grafico["reciclavel"][mes] = p_rec or 0.0
        dados_grafico["quimico"][mes] = p_qui or 0.0

    return render_template_string(
        HTML_FORMULARIO, 
        usuario_atual=usuario_logado, 
        e_admin=e_admin, 
        dados_grafico=dados_grafico
    )

@app.route("/salvar", methods=["POST"])
def salvar():
    if "usuario" not in session: return redirect(url_for("login"))
    
    data_str = request.form.get("data")
    mes_escolhido = request.form.get("mes")
    ano_escolhido = str(request.form.get("ano")).strip()
    localidade = request.form.get("localidade")
    
    qtd_inf = int(request.form.get("qtd_infectante", 0))
    peso_inf = float(request.form.get("peso_infectante", 0.0))
    qtd_org = int(request.form.get("qtd_organico", 0))
    peso_org = float(request.form.get("peso_organico", 0.0))
    qtd_comum = int(request.form.get("qtd_comum", 0))
    peso_comum = float(request.form.get("peso_comum", 0.0))
    qtd_rec = int(request.form.get("qtd_reciclavel", 0))
    peso_rec = float(request.form.get("peso_reciclavel", 0.0))
    qtd_qui = int(request.form.get("qtd_quimico", 0))
    peso_qui = float(request.form.get("peso_quimico", 0.0))

    dt = datetime.strptime(data_str, "%Y-%m-%d")
    data_formatada = dt.strftime("%d/%m/%Y")

    garantir_banco()
    conn = sqlite3.connect(CAMINHO_DB)
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO residuos (ano, mes, data, localidade, qtd_infectante, peso_infectante, qtd_organico, peso_organico, qtd_comum, peso_comum, qtd_reciclavel, peso_reciclavel, qtd_quimico, peso_quimico)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (ano_escolhido, mes_escolhido, data_formatada, localidade.upper(), qtd_inf, peso_inf, qtd_org, peso_org, qtd_comum, peso_comum, qtd_rec, peso_rec, qtd_qui, peso_qui))
    conn.commit()
    conn.close()

    flash(f"Lançamento salvo com sucesso ({mes_escolhido}/{ano_escolhido})!", "success")
    return redirect(url_for("index"))

@app.route("/importar/pdf", methods=["POST"])
def importar_pdf_sistema():
    if "usuario" not in session: return redirect(url_for("login"))
    
    if 'arquivo_pdf' not in request.files:
        flash("Nenhum ficheiro selecionado.", "error")
        return redirect(url_for("index"))
    
    file = request.files['arquivo_pdf']
    if file.filename == '':
        flash("Nome do ficheiro inválido.", "error")
        return redirect(url_for("index"))
    
    if file and file.filename.lower().endswith('.pdf'):
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as temp:
            file.save(temp.name)
            temp_path = temp.name

        try:
            reader = PdfReader(temp_path)
            garantir_banco()
            conn = sqlite3.connect(CAMINHO_DB)
            cursor = conn.cursor()

            meses_map = {
                "01": "JANEIRO", "02": "FEVEREIRO", "03": "MARÇO", "04": "ABRIL",
                "05": "MAIO", "06": "JUNHO", "07": "JULHO", "08": "AGOSTO",
                "09": "SETEMBRO", "10": "OUTUBRO", "11": "NOVEMBRO", "12": "DEZEMBRO"
            }

            total_importados = 0

            for page in reader.pages:
                texto = page.extract_text()
                if not texto: continue
                
                linhas = texto.split("\n")
                for linha in linhas:
                    m_data = re.search(r'(\d{2}/\d{2}/\d{4})', linha)
                    if m_data and "kg" in linha.lower():
                        data_str = m_data.group(1)
                        try:
                            dt_obj = datetime.strptime(data_str, "%d/%m/%Y")
                            ano = str(dt_obj.year)
                            mes_num = f"{dt_obj.month:02d}"
                            mes_nome = meses_map.get(mes_num, "OUTUBRO")
                            data_fmt = dt_obj.strftime("%d/%m/%Y")
                        except:
                            ano = "2026"
                            mes_nome = "OUTUBRO"
                            data_fmt = data_str

                        l_up = linha.upper()
                        setor_encontrado = "CENTRAL"
                        if "ANEXO H" in l_up: setor_encontrado = "ANEXO - H"
                        elif "ANEXO B" in l_up: setor_encontrado = "ANEXO - B"
                        elif "CENTRAL" in l_up or "PREDIO CENTRAL" in l_up: setor_encontrado = "CENTRAL"
                        elif "MATERNIDADE" in l_up: setor_encontrado = "MATERNIDADE"
                        elif "LABORATORIO" in l_up or "LABORATÓRIO" in l_up: setor_encontrado = "LABORATÓRIO"
                        elif "ALERGIA" in l_up: setor_encontrado = "CASA DA ALERGIA"
                        elif "NEFRO" in l_up: setor_encontrado = "CASA NEFRO"
                        elif "MAMOGRAFIA" in l_up: setor_encontrado = "MAMOGRAFIA"
                        elif "NUTRIÇÃO" in l_up or "NUTRICAO" in l_up: setor_encontrado = "NUTRIÇÃO"

                        match_pesos = re.findall(r'([\d.,]+)\s*kg', linha, re.IGNORECASE)
                        if match_pesos:
                            peso_str = match_pesos[-1].replace('.', '').replace(',', '.')
                            try:
                                peso_val = float(peso_str)
                                if peso_val > 0:
                                    cursor.execute("""
                                        INSERT INTO residuos (ano, mes, data, localidade, qtd_infectante, peso_infectante, qtd_organico, peso_organico, qtd_comum, peso_comum, qtd_reciclavel, peso_reciclavel, qtd_quimico, peso_quimico)
                                        VALUES (?, ?, ?, ?, 0, 0.0, 0, 0.0, 1, ?, 0, 0.0, 0, 0.0)
                                    """, (ano, mes_nome, data_fmt, setor_encontrado, peso_val))
                                    total_importados += 1
                            except:
                                pass

            conn.commit()
            conn.close()
            os.unlink(temp_path)
            
            if total_importados > 0:
                flash(f"PDF importado com sucesso! {total_importados} registos adicionados.", "success")
            else:
                flash("O PDF foi lido, mas nenhum registo compatível foi extraído.", "error")
        except Exception as e:
            if os.path.exists(temp_path): os.unlink(temp_path)
            flash(f"Erro ao processar o PDF: {str(e)}", "error")
    else:
        flash("Apenas ficheiros PDF são permitidos.", "error")
        
    return redirect(url_for("index"))

@app.route("/relatorio/filtros")
def relatorio_filtros():
    if "usuario" not in session: return redirect(url_for("login"))
    return render_template_string(HTML_FILTROS_RELATORIO)

@app.route("/relatorio/excel/gerar", methods=["POST"])
def relatorio_excel_gerar():
    if "usuario" not in session: return redirect(url_for("login"))
    
    filtro_localidade = request.form.get("filtro_localidade")
    filtro_ano = request.form.get("filtro_ano", "TODOS").strip()

    garantir_banco()
    conn = sqlite3.connect(CAMINHO_DB)
    cursor = conn.cursor()

    query = """
        SELECT ano, mes, data, localidade, 
               qtd_infectante, peso_infectante, 
               qtd_organico, peso_organico, 
               qtd_comum, peso_comum, 
               qtd_reciclavel, peso_reciclavel, 
               qtd_quimico, peso_quimico 
        FROM residuos WHERE 1=1
    """
    parametros = []

    if filtro_localidade != "TODOS":
        query += " AND localidade = ?"
        parametros.append(filtro_localidade.upper())
    if filtro_ano.upper() != "TODOS":
        query += " AND ano = ?"
        parametros.append(filtro_ano)

    cursor.execute(query, parametros)
    dados = cursor.fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relatório Consolidado"

    font_titulo = Font(name='Inter', size=14, bold=True, color='1E293B')
    font_subtitulo = Font(name='Inter', size=9, italic=True, color='64748B')
    font_cabecalho = Font(name='Inter', size=9, bold=True, color='FFFFFF')
    fill_cabecalho = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    borda_fina = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'), top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

    ws['A1'] = "RELATÓRIO DE GERENCIAMENTO DE RESÍDUOS - HOSPITALAR"
    ws['A1'].font = font_titulo
    ws['A2'] = f"Setor: {filtro_localidade} | Ano: {filtro_ano}"
    ws['A2'].font = font_subtitulo

    cabecalhos = [
        "Ano", "Mês", "Data", "Setor", 
        "Qtd. Inf.", "Peso Inf. (Kg)", 
        "Qtd. Org.", "Peso Org. (Kg)", 
        "Qtd. Comum", "Peso Comum (Kg)", 
        "Qtd. Recic.", "Peso Recic. (Kg)", 
        "Qtd. Quím.", "Peso Quím. (Kg)"
    ]

    for col_idx, texto in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=4, column=col_idx, value=texto)
        cell.font = font_cabecalho
        cell.fill = fill_cabecalho
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = borda_fina

    linha_destino = 5
    for row_data in dados:
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=linha_destino, column=col_idx, value=val)
            cell.font = Font(name='Inter', size=9, color='334155')
            cell.border = borda_fina
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = '@'
            elif col_idx in [2, 3, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [5, 7, 9, 11, 13]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0'
            elif col_idx in [6, 8, 10, 12, 14]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0.00'
        linha_destino += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"relatorio_residuos_completo.xlsx")

@app.route("/usuarios")
def pagina_usuarios():
    if "usuario" not in session or not usuario_admin(): return redirect(url_for("index"))
    return render_template_string(HTML_USUARIOS, lista_usuarios=carregar_usuarios())

@app.route("/usuarios/criar", methods=["POST"])
def criar_usuario():
    if "usuario" not in session or not usuario_admin(): return redirect(url_for("index"))
    usuarios = carregar_usuarios()
    novo_user = request.form.get("novo_usuario").strip().lower()
    senha = request.form.get("nova_senha").strip()
    if novo_user and senha and novo_user not in usuarios:
        usuarios[novo_user] = {"hash": generate_password_hash(senha), "role": "operador"}
        salvar_usuarios(usuarios)
        flash(f"Usuário '{novo_user}' criado!", "success")
    return redirect(url_for("pagina_usuarios"))

@app.route("/usuarios/deletar", methods=["POST"])
def deletar_usuario():
    if "usuario" not in session or not usuario_admin(): return redirect(url_for("index"))
    usuarios = carregar_usuarios()
    user_del = request.form.get("usuario_deletar")
    if user_del in usuarios and user_del not in ['admin', 'suporte']:
        del usuarios[user_del]
        salvar_usuarios(usuarios)
        flash("Usuário excluído!", "success")
    return redirect(url_for("pagina_usuarios"))

@app.route("/registros")
def listar_registros():
    if "usuario" not in session or not usuario_admin(): return redirect(url_for("index"))
    garantir_banco()
    conn = sqlite3.connect(CAMINHO_DB)
    cursor = conn.cursor()
    cursor.execute("SELECT id, ano, mes, data, localidade FROM residuos ORDER BY id DESC")
    rows = cursor.fetchall()
    conn.close()
    registros = [{"id": r[0], "ano": r[1], "mes": r[2], "data": r[3], "localidade": r[4]} for r in rows]
    return render_template_string(HTML_REGISTROS, registros=registros)

@app.route("/registro/excluir", methods=["POST"])
def excluir_registro():
    if "usuario" not in session or not usuario_admin(): return redirect(url_for("index"))
    reg_id = request.form.get("registro_id")
    if reg_id and reg_id.isdigit():
        garantir_banco()
        conn = sqlite3.connect(CAMINHO_DB)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM residuos WHERE id = ?", (reg_id,))
        conn.commit()
        conn.close()
        flash("Registro removido!", "success")
    return redirect(url_for("listar_registros"))

from waitress import serve
# local de acesso mais banco em Waitress
if __name__ == "__main__":
    garantir_banco()
    print("Servidor rodando com Waitress em http://0.0.0.0:8080")
    serve(app, host="0.0.0.0", port=8080)
